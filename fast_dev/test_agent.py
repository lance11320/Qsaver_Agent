import json
import pytest
import agent

def test_saved_effective_endpoint_is_used(monkeypatch):
    for key in ('QSAVER_API_PROFILE','QSAVER_API_BASE_URL','DASHSCOPE_API_KEY','QSAVER_API_KEY'):
        monkeypatch.delenv(key,raising=False)
    monkeypatch.setattr(agent,'config',lambda:{'api_profiles':[{'model_name':'test-terra','base_url':'',
        'effective_base_url':'https://example.invalid/v1/chat/completions','api_key':'test-only'}]})
    c=agent.Client('test-terra')
    assert c.base=='https://example.invalid/v1' and c.key=='test-only'

def question(**kw):
    return dict(kind='question',number='1',section='',stem='测试题',options=['A.甲','B.乙'],answer='',explanation='',source_pages=[1],complete=True,**kw)

def test_invalid_source_rejected():
    obj={'items':[question()]}
    with pytest.raises(ValueError):
        agent.validate(obj,[2],[2])

def test_shared_section_answer_fallback_requires_unique_numbers():
    q=question(group='paper');q['section']='【1-2 为题组】'
    s=dict(q,kind='solution',section='',answer='TF')
    answers={('paper','','1'):[s]}
    assert agent.match_solutions(q,[q],answers)==[s]
    assert agent.match_solutions(q,[q,dict(q,section='other')],answers)==[]
    assert agent.match_solutions(q,[q],{('other','','1'):[s]})==[]

def test_embedded_options_recovered_without_conflicting_text_loss():
    q=question();q.update(stem='判断以下说法正误：\nA.甲\nB.乙\nC.丙\nD.丁',options=[])
    agent.normalize_embedded_options(q)
    assert q['options']==['甲','乙','丙','丁'] and q['stem']=='判断以下说法正误：'
    q.update(stem='题干\nA.甲\nB.乙',options=['甲','乙'])
    agent.normalize_embedded_options(q);assert q['stem']=='题干'
    q.update(stem='题干\nA.甲\nB.乙',options=['甲','另一项'])
    agent.normalize_embedded_options(q);assert 'B.乙' in q['stem']

def test_recover_cross_page_question_from_discarded_raw(tmp_path,monkeypatch):
    docdir=tmp_path/('h'*16);docdir.mkdir()
    agent.dump(docdir/'raw_0000_0.json',{'items':[dict(question(),number='2',source_pages=[2,3])]})
    items=[dict(question(),source_file='paper.pdf',source_hash='h'*32,group='paper')]
    docs=[{'file':'paper.pdf','group':'paper','total_pages':4,'plan':{'continuous_numbering':True,'last_question_number':2}}]
    monkeypatch.setattr(agent,'prepare',lambda *a,**k:[{'page':p,'text':'','images':[]} for p in k['page_numbers']])
    class Fake:
        key='secret';calls=0
        def chat(self,*a):self.calls+=1;return {'items':[dict(question(),number='2',source_pages=[2,3])]}
    c=Fake();result=agent.recover_missing_questions(items,docs,c,tmp_path,progress=lambda s:None)
    assert [q['number'] for q in result]==['1','2'] and c.calls==1
    assert docs[0]['recovered_question_numbers']==[2]
    agent.recover_missing_questions([items[0]],docs,c,tmp_path,progress=lambda s:None)
    assert c.calls==1

def test_wrong_ownership_rejected():
    assert agent.validate({'items':[question()]},[1,2],[2])['items']==[]

def test_ingest_idempotent_and_backup(tmp_path,monkeypatch):
    monkeypatch.setattr(agent,'ROOT',tmp_path)
    q=question(review_flags=[],page_images=[])
    q['answer']='A'
    result=tmp_path/'result.json'
    agent.dump(result,{'partial':False,'run_id':'x','questions':[q]})
    assert agent.ingest(result)['added']==1
    assert agent.ingest(result)['added']==0
    assert len(list((tmp_path/'backups').glob('*.json')))==1

def test_partial_does_not_write(tmp_path,monkeypatch):
    monkeypatch.setattr(agent,'ROOT',tmp_path)
    p=tmp_path/'result.json'
    agent.dump(p,{'partial':True,'questions':[]})
    with pytest.raises(ValueError):
        agent.ingest(p)
    assert not (tmp_path/'questions.json').exists()

def test_corrupt_database_not_overwritten(tmp_path,monkeypatch):
    monkeypatch.setattr(agent,'ROOT',tmp_path)
    (tmp_path/'questions.json').write_text('broken')
    p=tmp_path/'result.json'
    agent.dump(p,{'partial':False,'questions':[]})
    with pytest.raises(ValueError):
        agent.ingest(p)
    assert (tmp_path/'questions.json').read_text()=='broken'
    assert not (tmp_path/'questions.lock').exists()

def test_review_blocks_uncertain(tmp_path,monkeypatch):
    monkeypatch.setattr(agent,'ROOT',tmp_path)
    p=tmp_path/'result.json'
    agent.dump(p,{'partial':False,'run_id':'x','questions':[question(review_flags=['OCR 不确定'])]})
    assert agent.ingest(p)['added']==0

def test_group_does_not_mix_papers():
    assert agent.group_name('模拟1答案.pdf')==agent.group_name('模拟1.pdf')
    assert agent.group_name('模拟1.pdf')!=agent.group_name('模拟2.pdf')

def test_resume_and_cross_document_answers(tmp_path,monkeypatch):
    # This fixture models extraction only; text review has separate client/contract tests.
    import text_review
    monkeypatch.setattr(text_review,'reviewer',lambda *args:None)
    monkeypatch.setattr(agent,'ROOT',tmp_path)
    a,b=tmp_path/'模拟1.txt',tmp_path/'模拟1答案.txt'
    a.write_text('1.测试题',encoding='utf-8'); b.write_text('1.A',encoding='utf-8')
    calls=[]
    class Fake:
        def __init__(self,*a):self.base='test';self.key='secret';self.usage=[]
        def chat(self,prompt,images):
            if '进行文档规划' in prompt:
                return {'role':'unknown'}
            calls.append(prompt)
            q=question()
            if '1.A' in prompt:q.update(kind='solution',stem='',answer='A')
            return {'items':[q]}
    monkeypatch.setattr(agent,'Client',Fake)
    _,result=agent.run([a,b],progress=lambda s:None)
    assert result['questions'][0]['answer']=='A'
    agent.run([a,b],progress=lambda s:None)
    assert len(calls)==2

def test_docx_tables_and_images(tmp_path):
    from docx import Document
    from PIL import Image
    d=Document();d.add_paragraph('1.题干');t=d.add_table(rows=1,cols=2);t.cell(0,0).text='A.甲';t.cell(0,1).text='B.乙'
    im=tmp_path/'im.png';Image.new('RGB',(20,20),'red').save(im);d.add_picture(str(im))
    path=tmp_path/'test.docx';d.save(path)
    units=agent.prepare(path,tmp_path/'pages')
    assert 'A.甲' in units[0]['text'] and len(units[0]['images'])==1

def test_reconcile_overlap_and_resume(tmp_path):
    q=question(source_hash='h',source_file='x.pdf',group='x',page_images=[])
    second=dict(q,source_pages=[1,2])
    class Fake:
        calls=0
        def chat(self,*args):
            self.calls+=1
            return {'same_question':True,'items':[dict(second)]}
    c=Fake()
    result=agent.reconcile_overlaps([q,second],c,tmp_path)
    assert len(result)==1 and result[0]['source_pages']==[1,2]
    assert len(agent.reconcile_overlaps([q,second],c,tmp_path))==1
    assert c.calls==1

def test_reconcile_failure_preserves_candidates(tmp_path):
    q=question(source_hash='h',source_file='x.pdf',group='x',page_images=[])
    class Fake:
        def chat(self,*a):raise ValueError('unavailable')
    assert len(agent.reconcile_overlaps([q,dict(q)],Fake(),tmp_path))==2

def test_export_option_labels_not_duplicated():
    assert agent.clean_options(['A.甲','B、乙','ATP'])==['甲','乙','ATP']
    assert agent.fingerprint({'stem':'X','options':['A.甲']})==agent.fingerprint({'stem':'X','options':['甲']})

def test_review_preserves_original(tmp_path):
    p=tmp_path/'result.json'
    agent.dump(p,{'questions':[question(review_flags=['答案复核不一致'])]})
    edited={'stem':'核对题','options':['甲'],'answer':'A','explanation':''}
    out,result=agent.review_question(p,0,edited)
    assert out.name=='result_reviewed.json'
    assert result['questions'][0]['review_flags']==[]
    assert json.loads(p.read_text(encoding='utf-8'))['questions'][0]['review_flags']

def test_answer_audit_blocks_mismatch(tmp_path):
    image=tmp_path/'0001.png'
    q=question(review_flags=[],answer_source={'source_file':'x','source_pages':[1]})
    q['answer']='FTTF'
    class Fake:
        def chat(self,*a):return {'answers':[{'number':'1','answer':'FTFF'}]}
    agent.verify_answers([q],[{'source_file':'x','page_images':[str(image)]}],Fake(),tmp_path)
    assert '答案复核不一致' in q['review_flags']
